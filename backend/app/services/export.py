from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import EmployeeProfile

COLUMNS = [
    ("№ п/п", 6), ("ФИО", 30), ("Должность", 26), ("Телефон", 20), ("Email", 28),
    ("Подразделение", 30), ("ФИО руководителя", 30), ("Должность руководителя", 26),
    ("Телефон руководителя", 20), ("№ приказа", 14), ("Дата приказа", 14),
]


def _text(value: str | None) -> str | None:
    # Строка, начинающаяся с «=», в Excel станет формулой. Экранируем, чтобы данные не исполнялись.
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def export_employees_xlsx(employees: list[EmployeeProfile]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append([title for title, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.cell(row=1, column=index).font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for number, employee in enumerate(employees, start=1):
        sheet.append(
            [
                number,
                _text(employee.full_name),
                _text(employee.position),
                _text(employee.phone),
                _text(employee.email),
                _text(employee.department),
                _text(employee.manager_full_name),
                _text(employee.manager_position),
                _text(employee.manager_phone),
                _text(employee.order_number),
                employee.order_date,
            ]
        )
        sheet.cell(row=sheet.max_row, column=11).number_format = "DD.MM.YYYY"

    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
